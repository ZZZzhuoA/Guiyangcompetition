from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.text.paragraph import Paragraph


DOCX_PATH = Path("docs/final_paper_style_report.docx")
BACKUP_PATH = Path("docs/final_paper_style_report.before_param_merge.docx")


def iter_block_items(document: Document):
    body = document.element.body
    for child in body.iterchildren():
        tag = child.tag.split("}")[-1]
        if tag == "p":
            yield "p", child
        elif tag == "tbl":
            yield "tbl", child


def paragraph_from_element(document: Document, p_el) -> Paragraph:
    return Paragraph(p_el, document)


def copy_first_run_font(src_para: Paragraph, dst_para: Paragraph) -> None:
    if not src_para.runs or not dst_para.runs:
        return
    src = src_para.runs[0]
    dst = dst_para.runs[0]
    dst.font.name = src.font.name
    dst.font.size = src.font.size
    dst.font.bold = src.font.bold
    dst.font.italic = src.font.italic
    if src.font.color and src.font.color.rgb:
        dst.font.color.rgb = src.font.color.rgb
    if src._element.rPr is not None and src._element.rPr.rFonts is not None:
        east_asia = src._element.rPr.rFonts.get(qn("w:eastAsia"))
        if east_asia:
            dst._element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:eastAsia"), east_asia)


def insert_paragraph_after(anchor: Paragraph, text: str) -> Paragraph:
    new_p = OxmlElement("w:p")
    anchor._p.addnext(new_p)
    paragraph = Paragraph(new_p, anchor._parent)
    paragraph.style = anchor.style
    paragraph.paragraph_format.space_before = anchor.paragraph_format.space_before
    paragraph.paragraph_format.space_after = anchor.paragraph_format.space_after
    paragraph.paragraph_format.line_spacing = anchor.paragraph_format.line_spacing
    paragraph.add_run(text)
    copy_first_run_font(anchor, paragraph)
    return paragraph


def find_para(document: Document, prefix: str) -> Paragraph:
    for para in document.paragraphs:
        if para.text.strip().startswith(prefix):
            return para
    raise RuntimeError(f"paragraph not found: {prefix}")


def main() -> None:
    doc = Document(DOCX_PATH)

    if not BACKUP_PATH.exists():
        import shutil

        shutil.copy2(DOCX_PATH, BACKUP_PATH)

    p53_anchor = find_para(doc, "因此，模型训练目标不是简单的历史策略类别")
    insert_paragraph_after(
        p53_anchor,
        "其中 beta=0.05 是任务优先级参数，而不是由评委两个指标等权直接得到的评分权重。由于 M1 和 M2 均已归一化到 [0,1]，该设置表示只有当两个策略的拦截率差异较小时，效费比差异才可能改变排序；若拦截率差距较大，则效费比不会推翻拦截率优先原则。",
    )

    p54_anchor = find_para(doc, "选择较浅树和较大的叶节点最小样本数")
    insert_paragraph_after(
        p54_anchor,
        "该参数组来自轻量随机森林基准实验。考虑到训练样本规模有限，而输入同时包含原始特征、工程特征和策略编码，树数量、树深度和叶节点样本数均采用偏稳健设置，以避免模型记忆训练场景中的局部偶然规律。该基准在验证集上取得 top1 0.5263、top2 0.7456、soft match 0.6218、regret 0.0425、Pareto hit 0.6754，因此保留为最终融合模型的主收益预测器。",
    )

    p55_anchor = find_para(doc, "训练过程中，对 Pairwise 输入特征进行标准化")
    insert_paragraph_after(
        p55_anchor,
        "Pairwise 模型只用于学习策略对相对优劣，不承担完整收益值预测，因此参数设置偏向稳定和轻量。epochs=80 用于保证收敛，learning_rate=0.08 在标准化特征下训练稳定，L2=0.003 用于抑制策略交互特征带来的过拟合风险。该组参数在策略对纠偏实验中能够稳定缓解策略3过推荐，并提高 top2 hit 与 soft match。",
    )

    p56_anchor = find_para(doc, "融合权重 0.65/0.35 的含义是")
    insert_paragraph_after(
        p56_anchor,
        "融合权重和策略3惩罚项通过验证集消融与误差诊断共同确定。仅加入 Pairwise 且 alpha=0.35 时，top1 由 0.5263 提升到 0.5439，top2 由 0.7456 提升到 0.7982，soft match 由 0.6218 提升到 0.6610；进一步加入 penalty_3=0.05 后，策略3过推荐从 +16 降至 +11，top2 提升到 0.8070，soft match 提升到 0.6666。相比之下，alpha=0.55 或 penalty=0.10 虽能继续降低策略3推荐次数，但会损失 top1 和 soft match，说明纠偏过强会破坏随机森林已经学到的有效收益规律。因此最终采用 alpha=0.35 与 penalty_3=0.05 的轻度纠偏配置。",
    )

    blocks = list(iter_block_items(doc))
    start = end = None
    for idx, (kind, element) in enumerate(blocks):
        if kind != "p":
            continue
        text = paragraph_from_element(doc, element).text.strip()
        if text.startswith("5.9参数、权重与校准项确定") or text.startswith("5.9 参数、权重与校准项确定"):
            start = idx
        if start is not None and text.startswith("6 实验设计"):
            end = idx
            break

    if start is None or end is None:
        raise RuntimeError(f"Could not identify 5.9 removal range: start={start}, end={end}")

    for _, element in blocks[start:end]:
        element.getparent().remove(element)

    doc.save(DOCX_PATH)
    print(f"updated: {DOCX_PATH}")
    print(f"backup: {BACKUP_PATH}")


if __name__ == "__main__":
    main()
