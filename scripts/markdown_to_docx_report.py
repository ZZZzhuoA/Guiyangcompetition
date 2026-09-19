from __future__ import annotations

import re
import sys
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


def repair_mojibake_line(line: str) -> str:
    try:
        repaired = line.encode("gbk").decode("utf-8")
    except UnicodeError:
        return line
    # Use the repaired form only when it removes common mojibake markers.
    bad_markers = ("鎷", "闅", "鏃", "瑁", "鐜", "绛", "锛", "銆", "鈥", "€")
    if any(m in line for m in bad_markers):
        return repaired
    return line


def clean_inline(text: str) -> str:
    text = text.replace("<br>", "\n")
    text = re.sub(r"`([^`]+)`", r"\1", text)
    text = re.sub(r"\*\*([^*]+)\*\*", r"\1", text)
    text = re.sub(r"\*([^*]+)\*", r"\1", text)
    text = re.sub(r"\[([^\]]+)\]\([^)]+\)", r"\1", text)
    return text


def set_cell_shading(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    tc_pr.append(shd)


def set_cell_margins(table, top=80, start=120, bottom=80, end=120):
    tbl_pr = table._tbl.tblPr
    tbl_cell_mar = tbl_pr.first_child_found_in("w:tblCellMar")
    if tbl_cell_mar is None:
        tbl_cell_mar = OxmlElement("w:tblCellMar")
        tbl_pr.append(tbl_cell_mar)
    for m, v in [("top", top), ("start", start), ("bottom", bottom), ("end", end)]:
        node = tbl_cell_mar.find(qn(f"w:{m}"))
        if node is None:
            node = OxmlElement(f"w:{m}")
            tbl_cell_mar.append(node)
        node.set(qn("w:w"), str(v))
        node.set(qn("w:type"), "dxa")


def set_repeat_table_header(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def set_run_font(run, font_name: str = "Calibri", east_asia: str = "Microsoft YaHei", size: float | None = None) -> None:
    run.font.name = font_name
    run._element.rPr.rFonts.set(qn("w:eastAsia"), east_asia)
    if size is not None:
        run.font.size = Pt(size)


def style_document(doc: Document) -> None:
    section = doc.sections[0]
    section.top_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.right_margin = Inches(1)

    normal = doc.styles["Normal"]
    normal.font.name = "Calibri"
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    normal.font.size = Pt(11)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.10

    for name, size, color, before, after in [
        ("Heading 1", 16, RGBColor(46, 116, 181), 16, 8),
        ("Heading 2", 13, RGBColor(46, 116, 181), 12, 6),
        ("Heading 3", 12, RGBColor(31, 77, 120), 8, 4),
    ]:
        style = doc.styles[name]
        style.font.name = "Calibri"
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
        style.font.size = Pt(size)
        style.font.color.rgb = color
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)


def add_paragraph(doc: Document, text: str, style: str | None = None):
    p = doc.add_paragraph(style=style)
    run = p.add_run(clean_inline(text))
    set_run_font(run)
    return p


def add_table(doc: Document, rows: list[list[str]]) -> None:
    if not rows:
        return
    col_count = max(len(r) for r in rows)
    normalized = [r + [""] * (col_count - len(r)) for r in rows]
    table = doc.add_table(rows=len(normalized), cols=col_count)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.style = "Table Grid"
    table.autofit = True
    set_cell_margins(table)
    if normalized:
        set_repeat_table_header(table.rows[0])
    wide = col_count >= 8
    font_size = 7.5 if wide else 9
    for r_idx, row in enumerate(normalized):
        for c_idx, value in enumerate(row):
            cell = table.cell(r_idx, c_idx)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            if r_idx == 0:
                set_cell_shading(cell, "F2F4F7")
            para = cell.paragraphs[0]
            para.alignment = WD_ALIGN_PARAGRAPH.CENTER if len(value) <= 14 else WD_ALIGN_PARAGRAPH.LEFT
            run = para.add_run(clean_inline(value.strip()))
            set_run_font(run, size=font_size)
            if r_idx == 0:
                run.bold = True
    doc.add_paragraph()


def parse_table(lines: list[str], start: int) -> tuple[list[list[str]], int]:
    rows: list[list[str]] = []
    i = start
    while i < len(lines) and lines[i].strip().startswith("|") and lines[i].strip().endswith("|"):
        raw = lines[i].strip().strip("|")
        parts = [p.strip() for p in raw.split("|")]
        # Skip markdown separator rows.
        if not all(re.fullmatch(r":?-{3,}:?", p.replace(" ", "")) for p in parts):
            rows.append(parts)
        i += 1
    return rows, i


def convert(md_path: Path, out_path: Path) -> None:
    raw_lines = md_path.read_text(encoding="utf-8").splitlines()
    lines = [repair_mojibake_line(line) for line in raw_lines]

    doc = Document()
    style_document(doc)
    in_code = False
    code_buffer: list[str] = []
    i = 0
    first_title = True
    while i < len(lines):
        line = lines[i].rstrip()
        stripped = line.strip()
        if not stripped:
            i += 1
            continue
        if stripped.startswith("```"):
            if in_code:
                if code_buffer:
                    p = doc.add_paragraph()
                    for code_line in code_buffer:
                        run = p.add_run(code_line + "\n")
                        run.font.name = "Consolas"
                        run._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
                        run.font.size = Pt(9)
                code_buffer = []
                in_code = False
            else:
                in_code = True
            i += 1
            continue
        if in_code:
            code_buffer.append(line)
            i += 1
            continue
        if stripped.startswith("|") and stripped.endswith("|"):
            table_rows, next_i = parse_table(lines, i)
            add_table(doc, table_rows)
            i = next_i
            continue
        if stripped.startswith("# "):
            text = clean_inline(stripped[2:].strip())
            p = doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            run = p.add_run(text)
            set_run_font(run, size=18)
            run.bold = True
            run.font.color.rgb = RGBColor(31, 77, 120)
            if first_title:
                first_title = False
            i += 1
            continue
        if stripped.startswith("## "):
            add_paragraph(doc, stripped[3:].strip(), style="Heading 1")
            i += 1
            continue
        if stripped.startswith("### "):
            add_paragraph(doc, stripped[4:].strip(), style="Heading 2")
            i += 1
            continue
        if stripped.startswith("#### "):
            add_paragraph(doc, stripped[5:].strip(), style="Heading 3")
            i += 1
            continue
        if re.match(r"^\d+\.\s+", stripped):
            add_paragraph(doc, re.sub(r"^\d+\.\s+", "", stripped), style="List Number")
            i += 1
            continue
        if stripped.startswith("- "):
            add_paragraph(doc, stripped[2:].strip(), style="List Bullet")
            i += 1
            continue
        add_paragraph(doc, stripped)
        i += 1

    doc.save(out_path)


if __name__ == "__main__":
    if len(sys.argv) != 3:
        raise SystemExit("Usage: markdown_to_docx_report.py input.md output.docx")
    convert(Path(sys.argv[1]), Path(sys.argv[2]))
