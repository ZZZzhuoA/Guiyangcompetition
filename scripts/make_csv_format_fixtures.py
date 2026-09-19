"""Generate schema-only CSV fixtures, with no simulation or outcome model.

The template's first two rows contain column names and primitive types.
Only template extraction needs openpyxl; CSV validation uses the standard library.
"""

import argparse
import csv
import hashlib
import json
import math
import re
from pathlib import Path


def extract_templates(directory):
    from openpyxl import load_workbook

    result = []
    families = {}
    for path in sorted(directory.glob("*.xlsx")):
        match = re.fullmatch(r"(.+)_(\d+)", path.stem)
        if not match:
            raise ValueError(f"Expected a numbered template: {path.name}")
        workbook = load_workbook(path, read_only=True, data_only=True)
        try:
            rows = workbook.active.iter_rows(values_only=True)
            headers = list(next(rows))
            types = list(next(rows))
        finally:
            workbook.close()
        while headers and headers[-1] is None:
            headers.pop()
        types = types[:len(headers)]
        if not headers or any(not isinstance(h, str) or not h for h in headers):
            raise ValueError(f"Empty or non-text header: {path.name}")
        if len(set(headers)) != len(headers):
            raise ValueError(f"Duplicate header: {path.name}")
        if len(types) != len(headers):
            raise ValueError(f"Header/type mismatch: {path.name}")
        types = [str(value).strip().lower() for value in types]
        if set(types) - {"int", "double", "bool", "string"}:
            raise ValueError(f"Unsupported primitive types: {path.name}: {types}")
        family, part = match.groups()
        signature = (headers, types)
        if family in families and families[family] != signature:
            raise ValueError(f"Different schemas in one family: {family}")
        families[family] = signature
        result.append({
            "file": path.with_suffix(".csv").name,
            "source": path.name,
            "source_sha256": hashlib.sha256(path.read_bytes()).hexdigest(),
            "family": family,
            "part": int(part),
            "columns": [dict(name=h, type=t) for h, t in zip(headers, types)],
            "time_columns": [h for h in headers if h.lower().endswith("time")],
        })
    if not result:
        raise ValueError(f"No XLSX templates in {directory}")
    return result


def parse_value(raw, primitive):
    if raw == "":
        return None
    if primitive == "string":
        return raw
    if primitive == "bool":
        if raw not in {"True", "False"}:
            raise ValueError(f"Expected True or False, got {raw!r}")
        return raw == "True"
    if primitive == "int":
        if not re.fullmatch(r"[+-]?\d+", raw):
            raise ValueError(f"Expected an integer, got {raw!r}")
        return int(raw)
    if primitive == "double":
        value = float(raw)
        if not math.isfinite(value):
            raise ValueError(f"Expected a finite number, got {raw!r}")
        return value
    raise ValueError(f"Unsupported primitive: {primitive}")


def read_typed_csv(path, template, allow_extra=True, allow_blank=True):
    """Read by column name; preserve blank values as None, never False/zero."""
    columns = template["columns"]
    names = [c["name"] for c in columns]
    with path.open("r", encoding="utf-8-sig", newline="") as stream:
        reader = csv.DictReader(stream, strict=True)
        headers = reader.fieldnames
        if headers is None or len(headers) != len(set(headers)):
            raise ValueError(f"Missing or duplicate headers: {path}")
        if set(names) - set(headers):
            raise ValueError(f"Required columns missing: {set(names) - set(headers)}")
        extra = [h for h in headers if h not in names]
        if extra and not allow_extra:
            raise ValueError(f"Unexpected columns: {extra}")
        records = []
        for raw in reader:
            if None in raw or any(v is None for v in raw.values()):
                raise ValueError(f"Wrong column count at CSV line {reader.line_num}")
            try:
                parsed = {c["name"]: parse_value(raw[c["name"]], c["type"])
                          for c in columns}
            except ValueError as exc:
                raise ValueError(f"{path.name}, CSV line {reader.line_num}: {exc}") from exc
            if not allow_blank and any(v is None for v in parsed.values()):
                raise ValueError(f"Blank cell at CSV line {reader.line_num}")
            records.append(parsed)
    return headers, extra, records


def make_rows(template, case, ticks):
    if case not in {"basic", "delayed_flags"}:
        raise ValueError(f"Unsupported fixture case: {case}")
    timed = bool(template["time_columns"])
    # Generic serialization tokens only. No column-specific physical meaning.
    for offset in range(ticks if timed else 1):
        row = []
        for index, col in enumerate(template["columns"]):
            name, primitive = col["name"], col["type"]
            if name in template["time_columns"]:
                value = float(template["part"] * ticks + offset)
            elif primitive == "bool":
                value = (template["part"] >= 3) if case == "delayed_flags" else offset % 2 == 1
            elif primitive == "int":
                value = 1
            elif primitive == "double":
                value = (offset % 4) / 4.0
            else:
                value = f'TEST,{index},"value"'
            row.append(value)
        yield row


def verify_fixtures(root, manifest):
    details = []
    for case in manifest["cases"]:
        for template in manifest["templates"]:
            path = root / case / template["file"]
            headers, extra, records = read_typed_csv(path, template, allow_blank=False)
            if headers != [c["name"] for c in template["columns"]] or extra:
                raise ValueError(f"Fixture headers differ from template: {path}")
            expected = list(make_rows(template, case, manifest["ticks_per_file"]))
            actual = [[record[h] for h in headers] for record in records]
            if actual != expected:
                raise ValueError(f"CSV round-trip changed values: {path}")
            for key in template["time_columns"]:
                times = [r[key] for r in records]
                if any(b <= a for a, b in zip(times, times[1:])):
                    raise ValueError(f"Non-increasing fixture time: {path}: {key}")
            values = [v for record in actual for v in record]
            details.append({
                "file": f"{case}/{template['file']}", "rows": len(records),
                "columns": len(headers), "blank_cells": sum(v is None for v in values),
                "true_cells": sum(v is True for v in values),
                "false_cells": sum(v is False for v in values),
                "sha256": hashlib.sha256(path.read_bytes()).hexdigest(),
            })
    for case in manifest["cases"]:
        by_family = {}
        for t in manifest["templates"]:
            by_family.setdefault(t["family"], []).append(t)
        for templates in by_family.values():
            for previous, current in zip(sorted(templates, key=lambda t: t["part"]),
                                         sorted(templates, key=lambda t: t["part"])[1:]):
                if current["part"] != previous["part"] + 1:
                    raise ValueError("Template file numbering contains a gap")
    return {"scope": "synthetic CSV format round-trip only", "passed": True,
            "file_count": len(details), "row_count": sum(d["rows"] for d in details),
            "files": details}


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--templates", type=Path)
    parser.add_argument("--output", type=Path, required=True)
    parser.add_argument("--ticks-per-file", type=int, default=200)
    parser.add_argument("--validate-only", action="store_true")
    args = parser.parse_args()
    if args.validate_only:
        manifest = json.loads((args.output / "manifest.json").read_text(encoding="utf-8"))
    else:
        if args.templates is None or args.ticks_per_file < 1:
            parser.error("Generation requires --templates and positive --ticks-per-file")
        if args.output.exists() and any(args.output.iterdir()):
            parser.error("Output must be empty; existing files will not be overwritten")
        templates = extract_templates(args.templates)
        manifest = {
            "synthetic": True, "purpose": "CSV parser fixtures, not training data",
            "encoding": "utf-8-sig", "delimiter": ",", "line_ending": "CRLF",
            "bool_tokens": ["True", "False"], "blank_token": "",
            "serialization_status": "Local test choices, not organizer-confirmed",
            "nullability_status": "Not tested; fixtures contain no blank values",
            "semantic_validation": False, "ticks_per_file": args.ticks_per_file,
            "cases": ["basic", "delayed_flags"],
            "templates": templates,
        }
        args.output.mkdir(parents=True, exist_ok=True)
        for case in manifest["cases"]:
            folder = args.output / case
            folder.mkdir()
            for template in templates:
                with (folder / template["file"]).open("w", encoding="utf-8-sig", newline="") as stream:
                    writer = csv.writer(stream, lineterminator="\r\n")
                    writer.writerow([c["name"] for c in template["columns"]])
                    writer.writerows(make_rows(template, case, args.ticks_per_file))
        (args.output / "manifest.json").write_text(
            json.dumps(manifest, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    report = verify_fixtures(args.output, manifest)
    if not args.validate_only:
        (args.output / "validation_report.json").write_text(
            json.dumps(report, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    print(json.dumps({k: v for k, v in report.items() if k != "files"}, ensure_ascii=True))


if __name__ == "__main__":
    main()
