from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.oxml import OxmlElement
from docx.text.paragraph import Paragraph
from docx.table import Table


DOCX_PATH = Path("docs/final_paper_style_report.docx")
BACKUP_PATH = Path("docs/final_paper_style_report.before_71_simplify.docx")
FALLBACK_PATH = Path("docs/final_paper_style_report_7_1_simplified.docx")


def iter_block_items(document: Document):
    body = document.element.body
    for child in body.iterchildren():
        tag = child.tag.split("}")[-1]
        if tag == "p":
            yield "p", child
        elif tag == "tbl":
            yield "tbl", child


def as_paragraph(document: Document, element) -> Paragraph:
    return Paragraph(element, document)


def as_table(document: Document, element) -> Table:
    return Table(element, document)


def replace_para(document: Document, prefix: str, text: str) -> None:
    for para in document.paragraphs:
        if para.text.strip().startswith(prefix):
            if para.runs:
                para.runs[0].text = text
                for run in para.runs[1:]:
                    run.text = ""
            else:
                para.add_run(text)
            return
    raise RuntimeError(f"paragraph not found: {prefix}")


def insert_paragraph_after_element(document: Document, element, text: str, style_name: str = "Normal") -> Paragraph:
    new_p = OxmlElement("w:p")
    element.addnext(new_p)
    paragraph = Paragraph(new_p, document._body)
    paragraph.style = style_name
    paragraph.add_run(text)
    return paragraph


def main() -> None:
    doc = Document(DOCX_PATH)

    if not BACKUP_PATH.exists():
        import shutil

        shutil.copy2(DOCX_PATH, BACKUP_PATH)

    replace_para(
        doc,
        "本文首先比较了直接分类、效用回归、聚类自适应",
        "本文首先从五类技术路线开展基线比较：一是直接学习历史策略标签的分类方法，二是预测候选策略收益的回归方法，三是按场景结构自适应选择策略的聚类方法，四是同时考虑两个指标的多目标方法，五是学习策略相对优劣的排序方法。为避免把实验部分写成模型清单，正文仅讨论方法类别及其结论，具体代表模型和缩写说明放在表注中。需要说明的是，直接分类方法主要评价历史策略标签预测能力，因此记录的是 accuracy 和 macro-F1；对能够产生候选策略得分或排序的方法，本文统一计算 top1、top2、soft match、regret 和 Pareto hit 等细粒度指标。",
    )

    replace_para(
        doc,
        "结果显示，核心25特征并不是对所有基线模型都带来提升",
        "结果显示，核心25特征并不是对所有基线模型都带来提升。均值效用模型不依赖场景特征，因此结果完全不变；线性模型也几乎没有变化，说明简单线性结构难以充分利用这些机理特征；核方法和单独排序模型在加入特征后略有下降，可能是因为维度增加后距离度量或线性偏好边界受到噪声影响。",
    )

    replace_para(
        doc,
        "提升最明显的是树模型相关方法",
        "提升最明显的是树模型相关方法。随机森林效用回归在加入核心25特征后，top1 从0.3897提升到0.4256，soft match 从0.6483提升到0.6836，regret 从0.0532降至0.0516，Pareto hit 从0.5590提升到0.5846。多目标树模型也从 top1 0.3487 提升到0.3692，regret 从0.0796降至0.0712。这说明核心25特征主要通过树模型的非线性划分能力发挥作用，能够帮助模型识别“高威胁集中 + 装备能力不足”“目标数量大 + 起效装备少”“类型承压不均衡”等场景结构。",
    )

    table7 = doc.tables[7]
    name_map = {
        "mean_utility_by_strategy": "均值效用",
        "ridge_utility_a1": "线性回归",
        "rbf_kernel_g0.2_a0.05": "核回归",
        "reg_forest_t25_d6": "随机森林",
        "cluster_local_k3": "局部聚类",
        "pairwise_logistic_e800": "Pairwise排序",
        "pareto_dual_metric_reg_forest": "多目标树模型",
    }
    for row in table7.rows[1:]:
        key = row.cells[0].text.strip()
        if key in name_map:
            row.cells[0].text = name_map[key]

    # Add a compact note immediately after the raw42/core25 comparison table.
    blocks = list(iter_block_items(doc))
    table_seen = 0
    note_exists = any("表注：表中线性回归对应 Ridge" in p.text for p in doc.paragraphs)
    if not note_exists:
        for kind, element in blocks:
            if kind == "tbl":
                if table_seen == 7:
                    insert_paragraph_after_element(
                        doc,
                        element,
                        "表注：表中线性回归对应 Ridge 效用回归，核回归对应 RBF 核效用回归，随机森林对应候选策略效用回归森林，局部聚类对应场景聚类后的局部策略选择，Pairwise排序对应策略对偏好学习，多目标树模型对应分别预测拦截率和效费比后的 Pareto 筛选方法。raw42 表示仅使用原始42维特征，core25 表示在 raw42 基础上加入筛选后的25个核心工程特征。",
                    )
                    break
                table_seen += 1

    try:
        doc.save(DOCX_PATH)
        print(f"updated: {DOCX_PATH}")
    except PermissionError:
        doc.save(FALLBACK_PATH)
        print(f"updated: {FALLBACK_PATH}")
        print(f"original locked: {DOCX_PATH}")
    print(f"backup: {BACKUP_PATH}")


if __name__ == "__main__":
    main()
